"""AI Hub 71886 → Supabase 적재 (rec_cases · rec_efficacy).

설계: docs/design/02-recommendations-data-spec.md §3·§4.
소유: 김민경. 임베딩 칸은 NULL로 두고 서지우가 채운다.

원본 zip은 읽기만 한다. 재실행해도 안전(멱등) — case_id / (inci,name_kr) upsert.

    uv run python -m scripts.load_recommendations_data --dry-run   # DB 없이 파싱·리포트만
    uv run python -m scripts.load_recommendations_data             # 실제 적재 (SUPABASE_* 필요)

환경: SUPABASE_URL, SUPABASE_SERVICE_ROLE_KEY (적재 시). 원본 경로는 --src 또는
COSMOS_AIHUB_SRC 로 지정, 기본값은 아래 DEFAULT_SRC.
"""

from __future__ import annotations

import argparse
import json
import logging
import os
import re
import sys
import tempfile
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.common.skin_concerns import CONCERN_LABELS

logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
log = logging.getLogger("load_rec")

DEFAULT_SRC = Path(
    os.environ.get(
        "COSMOS_AIHUB_SRC",
        str(Path.home() / "mkim/실무/data/03.스킨케어 성분-효능 추천 데이터/3.개방데이터"),
    )
)
TL_SUBDIR = "2.데이터(NIA)/Training/02.라벨링데이터"
OTHER_ZIP_SUBPATH = "1.데이터/Other/Other.zip"

# xlsx 헤더(공백·개행 제거 정규화) → 내부 필드명. PoC 9개 + 설계 요구 5개(권장피부타입·
# 사용상주의사항·권장농도·배합규제·참고문헌).  화학물성 5종은 임베딩 제외·페이로드 보존.
COLUMN_MAP: dict[str, str] = {
    "성분명(INCI)": "inci",
    "한글명": "name_kr",
    "효능": "efficacy",
    "제품적특성": "product_traits",
    "권장피부타입": "recommended_skin_types",
    "사용상주의사항(안전성)": "safety_note",
    "권장농도": "recommended_concentration",
    "배합규제": "regulation_note",
    "원시데이터출처(참고문헌file명)": "reference_source",
    "화학적물성": "properties",
    "용해도": "solubility",
    "분자식": "formula",
    "분자량": "weight",
    "원료출처": "source",
}

# 근거 식별자만 남긴다. URL·원문 인용 통째는 형식 불일치로 버려지고, 아래는 형식은
# 맞지만 값이 가짜인 것(원본의 30%가 '10.xxxx/xxxxx'·'12345678' 같은 템플릿)을 거른다.
# 사용자에게 임상 근거로 노출되므로 위양성(진짜를 버림)보다 위음성(가짜 노출)이 더 해롭다.
_EVIDENCE_RE = re.compile(r"^(PMID|PMCID|DOI)\s*[:\s]\s*(\S.*)$", re.IGNORECASE)
_PLACEHOLDER_RE = re.compile(r"x{3,}|1234567|placeholder", re.IGNORECASE)


def _norm_header(value: object) -> str:
    """헤더 셀의 모든 공백·개행을 제거해 매칭 키로 만든다."""
    return "".join(str(value).split()) if value is not None else ""


def _load_dotenv(path: Path = Path(".env")) -> None:
    """의존성 없이 .env의 KEY=VALUE를 os.environ에 채운다(기존 환경변수 우선)."""
    if not path.exists():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))


def clean_evidence(sources: list[str] | None) -> list[str]:
    """evidence_sources 정제 → 'PMID:xxxx' / 'DOI:xxxx' 형태만, 중복 제거."""
    out: list[str] = []
    for raw in sources or []:
        m = _EVIDENCE_RE.match(raw.strip())
        if not m:
            continue
        kind, val = m.group(1).upper(), m.group(2).strip()
        if _PLACEHOLDER_RE.search(val) or val.upper() in {"N/A", "NA"}:
            continue
        token = f"{kind}:{val}"
        if token not in out:
            out.append(token)
    return out


def build_name_matcher(ingredients: list[dict[str, str]]) -> re.Pattern[str] | None:
    """성분명 사전(한글명+INCI)으로 상담문에서 성분을 뽑을 정규식. 긴 이름 우선."""
    names = {n for ing in ingredients for n in (ing["name_kr"], ing["inci"]) if n and len(n) >= 2}
    if not names:
        return None
    ordered = sorted(names, key=len, reverse=True)
    return re.compile("|".join(re.escape(n) for n in ordered))


def _cot_step2_text(cot: list[dict[str, Any]]) -> str:
    return " ".join(c.get("content", "") for c in cot if c.get("step") == 2)


def extract_recommended(matcher: re.Pattern[str] | None, answer: str, cot: list) -> list[str]:
    """answer + CoT step2(성분 선택 근거)에서 사전 매칭된 성분명, 중복 제거."""
    if matcher is None:
        return []
    text = f"{answer}\n{_cot_step2_text(cot)}"
    seen: list[str] = []
    for m in matcher.finditer(text):
        if m.group(0) not in seen:
            seen.append(m.group(0))
    return seen


def read_cases(src: Path, matcher: re.Pattern[str] | None, limit: int | None) -> list[dict]:
    tl_dir = src / TL_SUBDIR
    zips = sorted(tl_dir.glob("TL_*.zip"))
    if not zips:
        log.error("TL zip 없음: %s", tl_dir)
        sys.exit(1)
    cases: list[dict] = []
    seen: set[str] = set()
    for zp in zips:
        with zipfile.ZipFile(zp) as zf:
            for name in zf.namelist():
                if not name.endswith(".jsonl"):
                    continue
                for line in zf.read(name).decode("utf-8").splitlines():
                    if not line.strip():
                        continue
                    rec = json.loads(line)
                    info, meta = rec.get("info", {}), rec.get("meta", {})
                    if not (info.get("id") and info.get("question") and info.get("answer")):
                        continue
                    if info["id"] in seen:
                        continue
                    seen.add(info["id"])
                    cot = rec.get("chain_of_thought", [])
                    cases.append(
                        {
                            "case_id": info["id"],
                            "target_concern": info.get("target_concern", ""),
                            "skin_concerns": meta.get("skin_concerns") or [],
                            "question": info["question"],
                            "answer": info["answer"],
                            "cot": cot,
                            "recommended_ingredients": extract_recommended(
                                matcher, info["answer"], cot
                            ),
                            "evidence_sources": clean_evidence(info.get("evidence_sources")),
                            "gender": meta.get("gender") or None,
                            "age": meta.get("age") or None,
                            "skin_type": meta.get("skin_type") or None,
                            "initial_skin_condition": meta.get("initial_skin_condition") or None,
                        }
                    )
                    if limit and len(cases) >= limit:
                        return cases
    return cases


def read_ingredients(src: Path) -> list[dict[str, str]]:
    with zipfile.ZipFile(src / OTHER_ZIP_SUBPATH) as zf:
        xlsx_name = next(n for n in zf.namelist() if n.endswith(".xlsx"))
        with tempfile.TemporaryDirectory() as td:
            path = zf.extract(xlsx_name, td)
            ws = load_workbook(path, read_only=True).worksheets[0]
            rows = ws.iter_rows(values_only=True)
            next(rows)  # 1행: 빈 행
            header = [_norm_header(c) for c in next(rows)]  # 2행: 헤더
            ings: list[dict[str, str]] = []
            for row in rows:
                raw = {h: v for h, v in zip(header, row, strict=False) if h}
                ing = {f: str(raw.get(col) or "").strip() for col, f in COLUMN_MAP.items()}
                if not ing["inci"] and not ing["name_kr"]:
                    continue
                ings.append(ing)
            return ings


def map_ingredient_ids(sb: Any, ings: list[dict]) -> int:
    """rec_efficacy 행에 ingredient_id 부여 (name_kr 정확 일치 → synonyms). 매칭 수 반환."""
    index: dict[str, int] = {}
    for table, keys in (("ingredients", ("name_kr",)), ("synonyms", ("name_kor", "synonym"))):
        start = 0
        while True:
            cols = "ingredient_id," + ",".join(keys)
            rows = sb.table(table).select(cols).range(start, start + 999).execute().data
            for r in rows:
                iid = r.get("ingredient_id")
                if not iid:
                    continue
                for k in keys:
                    if r.get(k):
                        index.setdefault(r[k], iid)
            if len(rows) < 1000:
                break
            start += 1000
    matched = 0
    for ing in ings:
        iid = index.get(ing["name_kr"]) or index.get(ing["inci"])
        ing["ingredient_id"] = iid
        matched += iid is not None
    return matched


def upsert(sb: Any, table: str, rows: list[dict], conflict: str, size: int = 500) -> None:
    for i in range(0, len(rows), size):
        sb.table(table).upsert(rows[i : i + size], on_conflict=conflict).execute()
        log.info("  %s upsert %d/%d", table, min(i + size, len(rows)), len(rows))


def report(cases: list[dict], ings: list[dict], id_matched: int | None) -> None:
    log.info("=== 적재 리포트 ===")
    log.info("rec_cases: %d건 / rec_efficacy: %d종", len(cases), len(ings))

    concern = Counter(c["target_concern"] for c in cases)
    log.info("고민(target_concern) 분포:")
    for label, n in concern.most_common():
        flag = "" if label in CONCERN_LABELS else "  <== CHECK 제약 밖!"
        log.info("  %5d  %s%s", n, label, flag)

    gender = Counter(c["gender"] for c in cases)
    ages = [c["age"] for c in cases if c["age"]]
    bucket: Counter[str] = Counter(f"{a // 10 * 10}대" for a in ages)
    log.info("성별 분포: %s", dict(gender))
    log.info("연령대 분포: %s", dict(sorted(bucket.items())))

    no_ing = sum(1 for c in cases if not c["recommended_ingredients"])
    no_ev = sum(1 for c in cases if not c["evidence_sources"])
    log.info("recommended_ingredients 0건 케이스: %d (%.1f%%)", no_ing, 100 * no_ing / len(cases))
    log.info("evidence_sources 0건 케이스(정제 후): %d (%.1f%%)", no_ev, 100 * no_ev / len(cases))
    ev_sample = [e for c in cases[:200] for e in c["evidence_sources"]][:8]
    log.info("정제된 근거 표본(PMID 스팟체크 대상): %s", ev_sample)

    if id_matched is not None:
        pct = 100 * id_matched / len(ings)
        log.info("ingredient_id 매칭: %d/%d (%.1f%%)", id_matched, len(ings), pct)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true", help="DB 없이 파싱·리포트만")
    ap.add_argument("--limit", type=int, default=None, help="케이스 상한(검증용)")
    ap.add_argument("--src", type=Path, default=DEFAULT_SRC, help="AI Hub 원본 루트")
    args = ap.parse_args()

    if not args.src.exists():
        log.error("원본 경로 없음: %s (--src 또는 COSMOS_AIHUB_SRC 지정)", args.src)
        sys.exit(1)

    ings = read_ingredients(args.src)
    matcher = build_name_matcher(ings)
    cases = read_cases(args.src, matcher, args.limit)
    # efficacy 는 임베딩 대상이자 NOT NULL. 빈 성분은 검색 신호가 없어 제외한다(정직 보고).
    empty_eff = sum(1 for i in ings if not i["efficacy"])
    ings = [i for i in ings if i["efficacy"]]
    log.info(
        "성분 지식: 원본 %d종 중 efficacy 빈 %d종 제외 → 적재 %d종",
        empty_eff + len(ings),
        empty_eff,
        len(ings),
    )

    if args.dry_run:
        report(cases, ings, id_matched=None)
        log.info("dry-run — DB 미적용")
        return

    from supabase import create_client  # 적재 시에만 필요

    _load_dotenv()
    url, key = os.environ.get("SUPABASE_URL"), os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
    if not (url and key):
        log.error("SUPABASE_URL·SUPABASE_SERVICE_ROLE_KEY 필요 (.env 또는 환경변수)")
        sys.exit(1)
    sb = create_client(url, key)

    id_matched = map_ingredient_ids(sb, ings)
    report(cases, ings, id_matched)
    log.info("적재 시작 …")
    upsert(sb, "rec_cases", cases, conflict="case_id")
    upsert(sb, "rec_efficacy", ings, conflict="inci,name_kr")
    log.info("적재 완료")


if __name__ == "__main__":
    main()
